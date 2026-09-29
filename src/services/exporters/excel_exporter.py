from pathlib import Path
from typing import Any, List

import duckdb
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.services.exporters.excel_styles import ExcelStyles


class ExportService:
    def __init__(
        self,
        conn: duckdb.DuckDBPyConnection,
        sql_file_path: Path,
        output_xlsx_path: Path,
        enterprise_name: str = "EMPRESA_DESCONHECIDA",
    ):
        self.conn = conn
        self.sql_file_path = sql_file_path
        self.output_xlsx_path = output_xlsx_path
        self.enterprise_name = (enterprise_name or "EMPRESA_DESCONHECIDA").upper()

    def export_query_to_excel(self) -> None:
        if not self.sql_file_path.exists():
            raise FileNotFoundError(f"Arquivo SQL não encontrado: {self.sql_file_path}")

        sql_query = self.sql_file_path.read_text(encoding="utf-8").strip()
        self.output_xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        cursor = self.conn.execute(sql_query)
        columns = [desc[0] for desc in cursor.description]

        wb = Workbook()
        ws = wb.active
        ws.title = "Auditoria Fiscal"
        ws.views.sheetView[0].showGridLines = True

        last_col_letter = get_column_letter(len(columns))

        self._build_header_section(ws, last_col_letter, columns)
        last_data_row = self._populate_data(ws, cursor, columns)

        self._apply_table_formatting(ws, last_col_letter, last_data_row)
        self._auto_fit_columns(ws)

        wb.save(self.output_xlsx_path)

    def _build_header_section(self, ws, last_col_letter: str, columns: List[str]) -> None:
        ws.merge_cells(f"A1:{last_col_letter}1")
        title_cell = ws["A1"]
        title_cell.value = "TRM Análises — RELATÓRIO DE AUDITORIA FISCAL"
        title_cell.font = ExcelStyles.TITLE_FONT
        title_cell.fill = ExcelStyles.TITLE_FILL
        title_cell.alignment = ExcelStyles.ALIGN_TITLE
        ws.row_dimensions[1].height = 36

        ws.merge_cells(f"A2:{last_col_letter}2")
        subtitle_cell = ws["A2"]
        subtitle_cell.value = f"EMPRESA AUDITADA: {self.enterprise_name}"
        subtitle_cell.font = ExcelStyles.SUBTITLE_FONT
        subtitle_cell.fill = ExcelStyles.SUBTITLE_FILL
        subtitle_cell.alignment = ExcelStyles.ALIGN_TITLE
        ws.row_dimensions[2].height = 24

        for col_idx, col_name in enumerate(columns, start=1):
            display_name = col_name.replace("_centavos", "").replace("_", " ").upper()
            cell = ws.cell(row=3, column=col_idx, value=display_name)
            cell.font = ExcelStyles.HEADER_FONT
            cell.fill = ExcelStyles.HEADER_FILL
            cell.border = ExcelStyles.THIN_BORDER
            cell.alignment = ExcelStyles.ALIGN_HEADER

        ws.row_dimensions[3].height = 28

    def _populate_data(self, ws, cursor: duckdb.DuckDBPyConnection, columns: List[str]) -> int:
        status_col_idx = columns.index("status_auditoria") if "status_auditoria" in columns else None
        current_row = 4

        while True:
            rows = cursor.fetchmany(1000)
            if not rows:
                break

            for row in rows:
                status_val = str(row[status_col_idx]) if status_col_idx is not None else ""
                row_fill, row_font = ExcelStyles.get_status_style(status_val)

                ws.row_dimensions[current_row].height = 22

                for col_idx, val in enumerate(row, start=1):
                    col_name = columns[col_idx - 1].lower()
                    cell = ws.cell(row=current_row, column=col_idx)

                    cell.border = ExcelStyles.THIN_BORDER
                    cell.font = row_font if row_font else ExcelStyles.BASE_FONT

                    if row_fill:
                        cell.fill = row_fill

                    self._format_cell_value(cell, col_name, val)

                current_row += 1

        return max(current_row - 1, 4)

    def _format_cell_value(self, cell, col_name: str, val: Any) -> None:
        if "modelo" in col_name:
            cell.value = ExcelStyles.format_nfe_model(val)
            cell.alignment = ExcelStyles.ALIGN_CENTER
        elif "valor" in col_name and isinstance(val, (int, float)):
            cell.value = val / 100.0 if isinstance(val, int) else val
            cell.number_format = "R$ #,##0.00"
            cell.alignment = ExcelStyles.ALIGN_RIGHT
        elif "data" in col_name:
            cell.value = val
            cell.alignment = ExcelStyles.ALIGN_CENTER
        elif isinstance(val, (int, float)):
            cell.value = val
            cell.number_format = "#,##0"
            cell.alignment = ExcelStyles.ALIGN_RIGHT
        else:
            cell.value = val
            cell.alignment = ExcelStyles.ALIGN_CENTER if col_name.startswith("status") else ExcelStyles.ALIGN_LEFT

    def _apply_table_formatting(self, ws, last_col_letter: str, last_data_row: int) -> None:
        """Converte o intervalo de dados em uma Tabela do Excel com estilo leve."""
        tab_range = f"A3:{last_col_letter}{last_data_row}"
        tab = Table(displayName="TabelaAuditoriaTRM", ref=tab_range)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(tab)
        ws.freeze_panes = "A4"

    def _auto_fit_columns(self, ws) -> None:
        """Ajusta a largura de cada coluna dinamicamente de acordo com o maior conteúdo."""
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.row in (1, 2):
                    continue
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)

            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 5, 15), 55)