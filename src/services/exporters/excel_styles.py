from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


class ExcelStyles:
    FONT_FAMILY = "Inter"

    TITLE_FONT = Font(name=FONT_FAMILY, size=13, bold=True, color="FFFFFF")
    SUBTITLE_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="2563EB")
    HEADER_FONT = Font(name=FONT_FAMILY, size=10, bold=True, color="FFFFFF")
    BASE_FONT = Font(name=FONT_FAMILY, size=10, color="0F172A")

    TITLE_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    SUBTITLE_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")

    THIN_BORDER = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
    ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
    ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")
    ALIGN_TITLE = Alignment(horizontal="left", vertical="center", indent=1)
    ALIGN_HEADER = Alignment(horizontal="center", vertical="center", wrap_text=True)

    STATUS_STYLES = {
        "XML_AUSENTE": {
            "fill": PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid"),
            "font": Font(name=FONT_FAMILY, size=10, bold=True, color="B91C1C"),
        },
        "WARNING": {
            "fill": PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"),
            "font": Font(name=FONT_FAMILY, size=10, bold=True, color="B45309"),
        },
        "OK": {
            "fill": PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid"),
            "font": Font(name=FONT_FAMILY, size=10, bold=True, color="15803D"),
        },
    }

    WARNING_STATUSES = {
        "SPED_AUSENTE",
        "DIVERGENCIA_VALOR",
        "DIVERGENCIA_STATUS",
        "DIVERGENCIA_DATA",
    }

    MODEL_MAP = {
        "65": "NFC-e",
        "55": "NF-e",
    }

    @classmethod
    def get_status_style(cls, status: str) -> tuple[PatternFill | None, Font | None]:
        if status == "XML_AUSENTE":
            style = cls.STATUS_STYLES["XML_AUSENTE"]
            return style["fill"], style["font"]

        if status in cls.WARNING_STATUSES:
            style = cls.STATUS_STYLES["WARNING"]
            return style["fill"], style["font"]

        if status == "OK":
            style = cls.STATUS_STYLES["OK"]
            return style["fill"], style["font"]

        return None, None

    @classmethod
    def format_nfe_model(cls, model_val: object) -> str:
        val_str = str(model_val).strip() if model_val is not None else ""
        return cls.MODEL_MAP.get(val_str, "Não identificado")