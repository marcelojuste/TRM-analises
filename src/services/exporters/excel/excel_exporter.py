from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from openpyxl import Workbook
from pathlib import Path
from typing import Any
import pandas as pd

from src.services.exporters.excel.excel_styles import ExcelStyles


class ExcelExport:
    def __init__(
        self,
        output_xlsx_path: Path,
        enterprise_name: str = "EMPRESA_DESCONHECIDA",
    ):
        self.output_xlsx_path = output_xlsx_path
        self.enterprise_name = (enterprise_name or "EMPRESA_DESCONHECIDA").upper()

    def export_dataframe_to_excel(self, df_results: pd.DataFrame) -> None:
        if df_results is None or df_results.empty:
            raise ValueError("O DataFrame de auditoria está vazio. Nenhum dado para exportar.")

        self.output_xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        columns = list(df_results.columns)

        wb = Workbook()
        ws = wb.active
        ws.title = "Auditoria Fiscal"
        ws.views.sheetView[0].showGridLines = True

        last_col_letter = get_column_letter(len(columns))

        self._build_header_section(ws, last_col_letter, columns)
        last_data_row = self._populate_data(ws, df_results, columns)

        self._apply_table_formatting(ws, last_col_letter, last_data_row)
        self._auto_fit_columns(ws)

        wb.save(self.output_xlsx_path)

    def _build_header_section(self, ws, last_col_letter: str, columns: list[str]) -> None:
        ws.merge_cells(f"A1:{last_col_letter}1")
        title_cell = ws["A1"]
        title_cell.value = "TRM Análises — RELATÓRIO DE AUDITORIA FISCAL"
        title_cell.font = ExcelStyles.TITLE_FONT
        title_cell.fill = ExcelStyles.TITLE_FILL
        title_cell.alignment = ExcelStyles.ALIGN_TITLE
        ws.row_dimensions[1].height = 36

        ws.merge_cells(f"A2:{last_col_letter}2")
        subtitle_cell = ws["A2"]
        subtitle_cell.value = f"EMPRESA AUDITADA: {self.enterprise_name}"
        subtitle_cell.font = ExcelStyles.SUBTITLE_FONT
        subtitle_cell.fill = ExcelStyles.SUBTITLE_FILL
        subtitle_cell.alignment = ExcelStyles.ALIGN_TITLE
        ws.row_dimensions[2].height = 24

        for col_idx, col_name in enumerate(columns, start=1):
            display_name = (
                col_name.replace("_cents", "")
                .replace("_centavos", "")
                .replace("_val", "")
                .replace("_", " ")
                .upper()
            )
            cell = ws.cell(row=3, column=col_idx, value=display_name)
            cell.font = ExcelStyles.HEADER_FONT
            cell.fill = ExcelStyles.HEADER_FILL
            cell.border = ExcelStyles.THIN_BORDER
            cell.alignment = ExcelStyles.ALIGN_HEADER

        ws.row_dimensions[3].height = 28

    def _populate_data(self, ws, df_results: pd.DataFrame, columns: list[str]) -> int:
        status_col_name = None
        for col in ("audit_status", "status_auditoria"):
            if col in columns:
                status_col_name = col
                break

        current_row = 4

        for row_tuple in df_results.itertuples(index=False):
            status_val = getattr(row_tuple, status_col_name, "") if status_col_name else ""
            row_fill, row_font = ExcelStyles.get_status_style(str(status_val))

            ws.row_dimensions[current_row].height = 22

            for col_idx, (col_name, val) in enumerate(zip(columns, row_tuple), start=1):
                cell = ws.cell(row=current_row, column=col_idx)

                cell.border = ExcelStyles.THIN_BORDER
                cell.font = row_font if row_font else ExcelStyles.BASE_FONT

                if row_fill:
                    cell.fill = row_fill

                self._format_cell_value(cell, col_name.lower(), val)

            current_row += 1

        return max(current_row - 1, 4)

    def _format_cell_value(self, cell, col_name: str, val: Any) -> None:
        if pd.isna(val) or val is None:
            cell.value = ""
            cell.alignment = ExcelStyles.ALIGN_CENTER
            return

        if "model" in col_name or "modelo" in col_name:
            cell.value = ExcelStyles.format_nfe_model(val)
            cell.alignment = ExcelStyles.ALIGN_CENTER
        elif ("value" in col_name or "valor" in col_name or "val" in col_name) and isinstance(
            val, (int, float)
        ):
            cell.value = val / 100.0 if "cents" in col_name or "centavos" in col_name else float(val)
            cell.number_format = "R$ #,##0.00"
            cell.alignment = ExcelStyles.ALIGN_RIGHT
        elif "date" in col_name or "data" in col_name:
            cell.value = str(val)
            cell.alignment = ExcelStyles.ALIGN_CENTER
        elif isinstance(val, (int, float)) and not isinstance(val, bool):
            cell.value = val
            cell.number_format = "#,##0"
            cell.alignment = ExcelStyles.ALIGN_RIGHT
        else:
            cell.value = str(val)
            cell.alignment = (
                ExcelStyles.ALIGN_CENTER if col_name.startswith("status") else ExcelStyles.ALIGN_LEFT
            )

    def _apply_table_formatting(self, ws, last_col_letter: str, last_data_row: int) -> None:
        tab_range = f"A3:{last_col_letter}{last_data_row}"
        tab = Table(displayName="TabelaAuditoriaTRM", ref=tab_range)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight1",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        ws.add_table(tab)
        ws.freeze_panes = "A4"

    def _auto_fit_columns(self, ws) -> None:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            
            for cell in col[2:]:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)

            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)
