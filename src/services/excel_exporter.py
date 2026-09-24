from pathlib import Path
import duckdb
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.writer.excel import WriteOnlyCell

class ExportService:
    @staticmethod
    def export_query_to_excel(
        conn: duckdb.DuckDBPyConnection, 
        sql_file_path: Path, 
        output_xlsx_path: Path
    ) -> None:
        if not sql_file_path.exists():
            raise FileNotFoundError(f"Arquivo SQL não encontrado: {sql_file_path}")

        sql_query = sql_file_path.read_text(encoding="utf-8").strip()
        output_xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        cursor = conn.execute(sql_query)
        columns = [desc[0] for desc in cursor.description]

        wb = Workbook(write_only=True)
        ws = wb.create_sheet(title="Dados Auditados")

        # Define os estilos desejados
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        cell_alignment = Alignment(vertical="center")

        header_row = []
        for col_name in columns:
            cell = WriteOnlyCell(ws, value=str(col_name).upper())
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = cell_alignment
            header_row.append(cell)
        
        ws.append(header_row)

        while True:
            rows = cursor.fetchmany(1000)
            if not rows:
                break
            
            for row in rows:
                formatted_row = []
                for val in row:
                    cell = WriteOnlyCell(ws, value=val)
                    cell.alignment = cell_alignment

                    if isinstance(val, float):
                        cell.number_format = '#,##0.00'
                    elif isinstance(val, int) and val > 1000:
                        cell.number_format = '#,##0'

                    formatted_row.append(cell)
                
                ws.append(formatted_row)

        wb.save(output_xlsx_path)