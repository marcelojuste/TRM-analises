import ctypes
from pathlib import Path
import sys
import threading
from tkinter import messagebox

import customtkinter as ctk
import openpyxl
from PIL import Image, ImageTk

from src.app_paths import PATHS
from src.exceptions import (
    AuditAppError,
    InterruptedException,
    NoSpedRecordsFoundException,
    NoXmlsFoundException,
)
from src.services.audit_service import AuditService
from src.views.components.file_card import FileSelectionCard
from src.views.components.header import HeaderFrame
from src.views.components.metric_card import MetricCard
from src.views.components.result_table import ResultsTableFrame

ctk.set_appearance_mode("Light")


class MainWindow(ctk.CTk):
    def __init__(self):
        super().__init__()

        self._set_windows_app_id()

        self.title("TRM Análises - Auditoria Fiscal")
        self.geometry("980x680")
        self.minsize(900, 650)
        self.configure(fg_color="#F8FAFC")

        self.is_running = False
        self.cancel_event = threading.Event()
        self.audit_thread = None

        self._set_app_icon()
        self._build_widgets()

    def _set_windows_app_id(self):
        if sys.platform.startswith("win"):
            my_app_id = "trmsistemas.trmanalises.auditor.1.0"
            try:
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(my_app_id)
            except Exception as e:
                print(f"Erro ao definir AppUserModelID: {e}")

    def _set_app_icon(self):
        icon_path = PATHS.icons_dir / "iconTRM.ico"
        if icon_path.exists():
            try:
                self.iconbitmap(str(icon_path))
            except Exception:
                img = Image.open(icon_path)
                photo = ImageTk.PhotoImage(img)
                self.iconphoto(False, photo)

    def _build_widgets(self):
        self.header = HeaderFrame(self)
        self.header.pack(fill="x", padx=30, pady=(24, 20))

        self.divider = ctk.CTkFrame(self, height=1, fg_color="#E2E8F0")
        self.divider.pack(fill="x", padx=30, pady=(0, 20))

        self.grid_container = ctk.CTkFrame(self, fg_color="transparent")
        self.grid_container.pack(fill="x", padx=30, pady=(0, 16))
        self.grid_container.grid_columnconfigure((0, 1), weight=1, uniform="card")

        self.card_nfce = FileSelectionCard(self.grid_container, title="NFC-e (XMLs)", is_directory=True)
        self.card_nfce.grid(row=0, column=0, padx=(0, 10), pady=(0, 15), sticky="ew")

        self.card_nfe = FileSelectionCard(self.grid_container, title="NF-e (XMLs)", is_directory=True)
        self.card_nfe.grid(row=0, column=1, padx=(10, 0), pady=(0, 15), sticky="ew")

        self.card_sped_fiscal = FileSelectionCard(self.grid_container, title="SPED Fiscal (.txt)", is_directory=False)
        self.card_sped_fiscal.grid(row=1, column=0, padx=(0, 10), pady=0, sticky="ew")

        self.card_sped_cofins = FileSelectionCard(self.grid_container, title="SPED Contribuições (.txt)", is_directory=False)
        self.card_sped_cofins.grid(row=1, column=1, padx=(10, 0), pady=0, sticky="ew")

        self.file_cards = [
            self.card_nfce, 
            self.card_nfe, 
            self.card_sped_fiscal, 
            self.card_sped_cofins
        ]

        self.metrics_container = ctk.CTkFrame(self, fg_color="transparent")
        self.metrics_container.pack(fill="x", padx=30, pady=(0, 16))
        self.metrics_container.grid_columnconfigure((0, 1, 2, 3), weight=1, uniform="metric")

        self.card_xml_qty = MetricCard(self.metrics_container, title="Qtd. XMLs", value="0", subtitle="NF-e + NFC-e")
        self.card_xml_qty.grid(row=0, column=0, padx=(0, 6), sticky="ew")

        self.card_xml_val = MetricCard(self.metrics_container, title="Total XMLs", value="R$ 0,00", subtitle="Soma dos XMLs")
        self.card_xml_val.grid(row=0, column=1, padx=6, sticky="ew")

        self.card_sped_qty = MetricCard(self.metrics_container, title="Linhas SPED", value="0", subtitle="Registros processados")
        self.card_sped_qty.grid(row=0, column=2, padx=6, sticky="ew")

        self.card_sped_val = MetricCard(self.metrics_container, title="Total SPED", value="R$ 0,00", subtitle="Soma do SPED")
        self.card_sped_val.grid(row=0, column=3, padx=(6, 0), sticky="ew")

        self.footer_frame = ctk.CTkFrame(self, fg_color="transparent")
        self.footer_frame.pack(fill="x", padx=30, pady=(0, 16))
        self.footer_frame.grid_columnconfigure((0, 1), weight=1, uniform="footer_btn")

        self.btn_execute = ctk.CTkButton(
            self.footer_frame,
            text="▶   Executar Auditoria",
            font=ctk.CTkFont(family="Inter", size=15, weight="bold"),
            fg_color="#0F172A",
            hover_color="#1E293B",
            text_color="#FFFFFF",
            corner_radius=8,
            height=46,
            command=self.run_audit
        )
        self.btn_execute.pack(fill="x", ipady=2)

        self.btn_clear = ctk.CTkButton(
            self.footer_frame,
            text="🔄   Limpar",
            font=ctk.CTkFont(family="Inter", size=15, weight="bold"),
            fg_color="#FEE2E2",
            hover_color="#FCA5A5",
            text_color="#991B1B",
            corner_radius=8,
            height=46,
            command=self._on_clear_or_cancel_clicked
        )

        self.results_table = ResultsTableFrame(self)

    def _set_ui_processing_state(self, is_processing: bool):
        self.is_running = is_processing
        enabled = not is_processing

        for card in self.file_cards:
            card.set_enabled(enabled)

        self.results_table.set_enabled(enabled)

        if is_processing:
            self.btn_execute.configure(state="disabled", text="⏳ Processando Auditoria...")
            
            self.btn_clear.configure(
                text="🛑   Cancelar Operação",
                fg_color="#EF4444",
                hover_color="#DC2626",
                text_color="#FFFFFF",
                state="normal"
            )

            if not self.btn_clear.winfo_ismapped():
                self.btn_execute.pack_forget()
                self.btn_execute.grid(row=0, column=0, padx=(0, 8), sticky="ew", ipady=2)
                self.btn_clear.grid(row=0, column=1, padx=(8, 0), sticky="ew", ipady=2)
        else:
            self.btn_execute.configure(state="normal", text="▶   Executar Auditoria")
            
            self.btn_clear.configure(
                text="🔄   Limpar",
                fg_color="#FEE2E2",
                hover_color="#FCA5A5",
                text_color="#991B1B",
                state="normal"
            )

    def _on_clear_or_cancel_clicked(self):
        if self.is_running:
            self.cancel_audit()
        else:
            self.reset_app()

    def cancel_audit(self):
        if self.is_running:
            self.cancel_event.set()
            self.btn_clear.configure(state="disabled", text="⏳ Cancelando...")

    def _cleanup_temp_files(self):
        try:
            if PATHS.outputs_dir.exists():
                for temp_file in PATHS.outputs_dir.glob("*.tmp"):
                    temp_file.unlink(missing_ok=True)
                for temp_file in PATHS.outputs_dir.glob("~$*.xlsx"):
                    temp_file.unlink(missing_ok=True)
        except Exception as e:
            print(f"Erro ao limpar arquivos temporários: {e}")

    def _update_metrics_cards(self, metrics: dict):
        xml_data = metrics.get("xml", {"qty": 0, "val": 0.0})
        sped_data = metrics.get("sped", {"qty": 0, "val": 0.0})

        formatted_xml_qty = f"{xml_data.get('qty', 0):,}".replace(",", ".")
        formatted_xml_val = f"R$ {xml_data.get('val', 0.0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        formatted_sped_qty = f"{sped_data.get('qty', 0):,}".replace(",", ".")
        formatted_sped_val = f"R$ {sped_data.get('val', 0.0):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

        self.card_xml_qty.set_value(formatted_xml_qty)
        self.card_xml_val.set_value(formatted_xml_val)
        self.card_sped_qty.set_value(formatted_sped_qty)
        self.card_sped_val.set_value(formatted_sped_val)

    def _read_generated_excel(self, excel_path: Path, max_rows: int = 100):
        wb = openpyxl.load_workbook(excel_path, read_only=False, data_only=True)
        sheet = wb.active

        columns = []
        data_rows = []

        header_row_idx = 3
        for col_idx in range(1, sheet.max_column + 1):
            val = sheet.cell(row=header_row_idx, column=col_idx).value
            columns.append(str(val) if val is not None else f"COLUNA_{col_idx}")

        start_data_row = 4
        for r_idx in range(start_data_row, min(sheet.max_row + 1, start_data_row + max_rows)):
            row_vals = []
            for c_idx in range(1, len(columns) + 1):
                cell = sheet.cell(row=r_idx, column=c_idx)
                val = cell.value

                if isinstance(val, (int, float)) and "valor" in str(columns[c_idx-1]).lower():
                    val = f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

                row_vals.append(val)

            if any(val is not None for val in row_vals):
                data_rows.append(row_vals)

        wb.close()
        return columns, data_rows

    def _show_results_table(self, excel_path: Path, company_name: str = None):
        if not excel_path.exists():
            return

        columns, rows = self._read_generated_excel(excel_path)
        if not columns:
            return

        self.results_table.load_excel_data(
            columns=columns, 
            rows=rows, 
            company_name=company_name, 
            excel_path=excel_path
        )
        self.results_table.pack(fill="both", expand=True, padx=30, pady=(0, 24))

        if self.winfo_height() < 800:
            self.geometry(f"{self.winfo_width()}x820")

    def run_audit(self):
        raw_paths = {
            "nfe": self.card_nfe.get_path(),
            "nfce": self.card_nfce.get_path(),
            "sped_fiscal": self.card_sped_fiscal.get_path(),
            "sped_cofins": self.card_sped_cofins.get_path(),
        }

        has_xml = any([raw_paths["nfe"], raw_paths["nfce"]])
        has_sped = any([raw_paths["sped_fiscal"], raw_paths["sped_cofins"]])

        if not has_xml and not has_sped:
            messagebox.showwarning(
                "Seleção Insuficiente", 
                "Nenhum arquivo ou diretório foi selecionado.\n\nPor favor, informe ao menos uma pasta de XMLs e um arquivo SPED para realizar a auditoria."
            )
            return

        if not has_xml:
            messagebox.showwarning(
                "XMLs Não Selecionados", 
                "Por favor, selecione ao menos um diretório de XMLs (NF-e ou NFC-e)."
            )
            return

        if not has_sped:
            messagebox.showwarning(
                "SPED Não Selecionado", 
                "Por favor, selecione ao menos um arquivo SPED (.txt)."
            )
            return

        paths = {key: Path(val) if val else None for key, val in raw_paths.items()}

        invalid_paths = [str(p) for p in paths.values() if p and not p.exists()]
        if invalid_paths:
            messagebox.showerror(
                "Caminho Inexistente", 
                "Os seguintes arquivos/diretórios não existem:\n\n" + "\n".join(f"• {p}" for p in invalid_paths)
            )
            return

        invalid_extensions = []
        for key in ["sped_fiscal", "sped_cofins"]:
            p = paths[key]
            if p and p.suffix.lower() != ".txt":
                invalid_extensions.append(f"{p.name} (esperado .txt)")

        if invalid_extensions:
            messagebox.showerror(
                "Formato Inválido",
                "Os arquivos do SPED devem ser arquivos de texto (.txt):\n\n" + "\n".join(f"• {item}" for item in invalid_extensions)
            )
            return

        self.cancel_event.clear()
        self._set_ui_processing_state(True)
        self.update_idletasks()

        def worker():
            cancelled = False
            try:
                audit_service = AuditService(
                    nfe_dir=paths["nfe"],
                    nfce_dir=paths["nfce"],
                    sped_fiscal_path=paths["sped_fiscal"],
                    sped_cofins_path=paths["sped_cofins"],
                    cancel_event=self.cancel_event
                )
                
                result = audit_service.run_pipeline()

                if self.cancel_event.is_set():
                    cancelled = True
                else:
                    excel_file = Path(result["excel_path"]) if result.get("excel_path") and Path(str(result["excel_path"])).exists() else None
                    enterprise_name = result.get("enterprise_name", "EMPRESA_DESCONHECIDA")
                    metrics = result.get("metrics", {})

                    self.after(0, lambda: self._update_metrics_cards(metrics))

                    if not excel_file:
                        generated_files = sorted(PATHS.outputs_dir.glob("*.xlsx"), key=lambda f: f.stat().st_mtime, reverse=True)
                        if generated_files:
                            excel_file = generated_files[0]

                    if excel_file:
                        self.after(0, lambda: self._show_results_table(excel_file, company_name=enterprise_name))

                    self.after(0, lambda: messagebox.showinfo(
                        "Sucesso", 
                        f"Auditoria concluída com sucesso!\n\nRelatório gerado em:\n{PATHS.outputs_dir.resolve()}"
                    ))

            except InterruptedException:
                cancelled = True
            except NoXmlsFoundException as nxe:
                self.after(0, lambda msg=str(nxe): messagebox.showwarning(
                    "XMLs Não Encontrados", 
                    f"Não foi possível encontrar os arquivos XML.\n\n{msg}"
                ))
            except NoSpedRecordsFoundException as nse:
                self.after(0, lambda msg=str(nse): messagebox.showwarning(
                    "Registros SPED Não Encontrados", 
                    f"Não foi possível encontrar os registros SPED.\n\n{msg}"
                ))
            except AuditAppError as aae:
                self.after(0, lambda msg=str(aae): messagebox.showerror(
                    "Erro da Aplicação", 
                    f"Ocorreu uma falha no fluxo de auditoria:\n\n{msg}"
                ))
            except PermissionError as pe:
                self.after(0, lambda msg=str(pe): messagebox.showerror("Arquivo em Uso", msg))
            except FileNotFoundError as fnf:
                self.after(0, lambda msg=str(fnf): messagebox.showerror("Arquivo Não Encontrado", f"Falha ao acessar o arquivo:\n{msg}"))
            except UnicodeDecodeError:
                self.after(0, lambda: messagebox.showerror(
                    "Erro de Codificação", 
                    "Falha ao ler um dos arquivos SPED. Verifique se o arquivo está salvo na codificação correta (UTF-8 ou ISO-8859-1)."
                ))
            except Exception as e:
                if self.cancel_event.is_set():
                    cancelled = True
                else:
                    error_msg = str(e)
                    self.after(0, lambda msg=error_msg: messagebox.showerror(
                        "Erro na Auditoria", 
                        f"Ocorreu um erro durante a execução:\n\n{msg}"
                    ))
            finally:
                if cancelled or self.cancel_event.is_set():
                    self._cleanup_temp_files()
                    self.after(0, self.reset_app)
                    self.after(0, lambda: messagebox.showwarning("Cancelado", "A operação de auditoria foi cancelada pelo usuário."))
                else:
                    self.after(0, lambda: self._set_ui_processing_state(False))

        self.audit_thread = threading.Thread(target=worker, daemon=True)
        self.audit_thread.start()

    def reset_app(self):
        for card in self.file_cards:
            card.clear_selection()

        self.card_xml_qty.reset("0", "NF-e + NFC-e")
        self.card_xml_val.reset("R$ 0,00", "Soma dos XMLs")
        self.card_sped_qty.reset("0", "Registros processados")
        self.card_sped_val.reset("R$ 0,00", "Soma do SPED")

        self.results_table.clear()
        self.results_table.pack_forget()

        self.btn_clear.grid_forget()
        self.btn_execute.grid_forget()
        self.btn_execute.pack(fill="x", ipady=2)

        self._set_ui_processing_state(False)
        self.geometry("980x680")


if __name__ == "__main__":
    app = MainWindow()
    app.mainloop()