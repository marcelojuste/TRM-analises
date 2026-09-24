from pathlib import Path
from unittest.mock import MagicMock, patch
import pytest
from openpyxl import load_workbook

from src.app_paths import AppPaths
from src.services.audit_service import AuditService


@pytest.fixture
def mock_paths(tmp_path, monkeypatch):
    root_dir = tmp_path
    database_dir = root_dir / "database"
    queries_dir = database_dir / "queries"
    temp_files_dir = root_dir / "temp_files"
    outputs_dir = root_dir / "outputs"
    schema_path = queries_dir / "schema.sql"

    queries_dir.mkdir(parents=True, exist_ok=True)
    temp_files_dir.mkdir(parents=True, exist_ok=True)
    outputs_dir.mkdir(parents=True, exist_ok=True)

    (queries_dir / "audit.sql").write_text("SELECT 1;", encoding="utf-8")
    schema_path.write_text("CREATE TABLE test (id INT);", encoding="utf-8")

    fake_paths = AppPaths(
        root_dir=root_dir,
        database_dir=database_dir,
        queries_dir=queries_dir,
        schema_path=schema_path,
        temp_files_dir=temp_files_dir,
        images_dir=root_dir / "images",
        icons_dir=root_dir / "icons",
        outputs_dir=outputs_dir,
    )

    monkeypatch.setattr("src.services.audit_service.PATHS", fake_paths)
    monkeypatch.setattr("src.database.database.PATHS", fake_paths)
    return fake_paths


def test_audit_service_initialization():
    xml_dir = Path("/tmp/xmls")
    sped_path = Path("/tmp/sped.txt")

    service = AuditService(xml_dir=xml_dir, sped_path=sped_path)

    assert service.xml_dir == xml_dir
    assert service.sped_path == sped_path


@patch("src.services.audit_service.ExportService.export_query_to_excel")
@patch("src.services.audit_service.FiscalRepository")
@patch("src.services.audit_service.DisposableAuditDatabase")
@patch("src.services.audit_service.XmlParser.get_xml_files")
@patch("src.services.audit_service.SpedParser")
def test_run_pipeline_flow_orchestration(
    mock_sped_parser,
    mock_get_xml_files,
    mock_db,
    mock_repo,
    mock_exporter,
    mock_paths,
):
    mock_sped_instance = MagicMock()
    mock_sped_instance.fiscal_notes = [("SPED_NOTE_1",)]
    mock_sped_instance.get_enterprise.return_value = "EMPRESA_TESTE_LTDA"
    mock_sped_parser.return_value.__enter__.return_value = mock_sped_instance

    mock_get_xml_files.return_value = []

    service = AuditService(xml_dir=Path("/fake/xmls"), sped_path=Path("/fake/sped.txt"))
    service.run_pipeline()

    mock_db.assert_called_once_with(enterprise="EMPRESA_TESTE_LTDA")

    repo_instance = mock_repo.return_value.__enter__.return_value
    repo_instance.add_sped.assert_called_once_with([("SPED_NOTE_1",)])

    expected_output_path = mock_paths.outputs_dir / "relatorio_auditoria_EMPRESA_TESTE_LTDA.xlsx"
    mock_exporter.assert_called_once_with(
        conn=mock_db.return_value.__enter__.return_value,
        sql_file_path=mock_paths.queries_dir / "audit.sql",
        output_xlsx_path=expected_output_path,
    )


@patch("src.services.audit_service.ExportService.export_query_to_excel")
@patch("src.services.audit_service.FiscalRepository")
@patch("src.services.audit_service.DisposableAuditDatabase")
@patch("src.services.audit_service.XmlParser")
@patch("src.services.audit_service.SpedParser")
def test_run_pipeline_fallback_enterprise_name(
    mock_sped_parser,
    mock_xml_parser,
    mock_db,
    mock_repo,
    mock_exporter,
    mock_paths,
):
    mock_sped_instance = MagicMock()
    mock_sped_instance.fiscal_notes = []
    mock_sped_instance.get_enterprise.return_value = None
    mock_sped_parser.return_value.__enter__.return_value = mock_sped_instance

    mock_xml_parser.get_xml_files.return_value = []

    service = AuditService(xml_dir=Path("/fake/xmls"), sped_path=Path("/fake/sped.txt"))
    service.run_pipeline()

    mock_db.assert_called_once_with(enterprise="EMPRESA_DESCONHECIDA")
    expected_output_path = mock_paths.outputs_dir / "relatorio_auditoria_EMPRESA_DESCONHECIDA.xlsx"
    assert mock_exporter.call_args[1]["output_xlsx_path"] == expected_output_path


@patch("src.services.audit_service.XmlParser")
def test_extract_xmls_success_and_ignore_corrupted(mock_xml_parser):
    fake_file_1 = Path("nfe1.xml")
    fake_file_2 = Path("corrupted.xml")

    mock_xml_parser.get_xml_files.return_value = [fake_file_1, fake_file_2]

    doc_valid = MagicMock()
    doc_valid.to_tuple.return_value = ("KEY123", 10050)

    mock_xml_parser.parse_xml.side_effect = [doc_valid, None]

    service = AuditService(xml_dir=Path("/fake/xmls"), sped_path=Path("/fake/sped.txt"))
    result = service._extract_xmls()

    assert len(result) == 1
    assert result[0] == ("KEY123", 10050)

@patch("src.services.audit_service.XmlParser")
@patch("src.services.audit_service.SpedParser")
def test_audit_service_real_pipeline_execution(
    mock_sped_parser, mock_xml_parser, mock_paths, tmp_path
):
    mock_sped_instance = MagicMock()
    mock_sped_instance.get_enterprise.return_value = "EMPRESA TESTE REAL SA"
    mock_sped_instance.fiscal_notes = [("C100", "12345")]
    mock_sped_parser.return_value.__enter__.return_value = mock_sped_instance

    fake_doc = MagicMock()
    fake_doc.to_tuple.return_value = ("NFe31240112345678000195550010000123451000123456", 10000)
    mock_xml_parser.get_xml_files.return_value = [tmp_path / "nfe.xml"]
    mock_xml_parser.parse_xml.return_value = fake_doc

    mock_paths.schema_path.write_text("""
        CREATE TABLE IF NOT EXISTS xml_documents (
            access_key VARCHAR,
            total_cents INT
        );
        CREATE TABLE IF NOT EXISTS sped_documents (
            reg VARCHAR,
            num_doc VARCHAR
        );
    """, encoding="utf-8")

    (mock_paths.queries_dir / "audit.sql").write_text("""
        SELECT access_key AS chave, total_cents AS valor FROM xml_documents
        UNION ALL
        SELECT reg AS chave, CAST(num_doc AS INT) AS valor FROM sped_documents;
    """, encoding="utf-8")

    service = AuditService(xml_dir=tmp_path, sped_path=tmp_path / "sped.txt")
    service.run_pipeline()

    expected_excel = mock_paths.outputs_dir / "relatorio_auditoria_EMPRESA TESTE REAL SA.xlsx"
    assert expected_excel.exists(), "O arquivo Excel não foi gerado no disco!"

    wb = load_workbook(expected_excel)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    assert len(rows) > 1, f"Relatório vazio! Linhas encontradas: {len(rows)}"
    assert len(rows) == 3